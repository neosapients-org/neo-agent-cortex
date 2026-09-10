"""Build an Excel report from flow_test_results.json.

Per-question layout so raw data + accuracy can be added by hand later:
  SrNo | Question | Run1/2/3 status | Platform Responded? | Consistent? |
  MCP hit | Avg latency (s) | Run1/2/3 response | Raw Data (blank) | Accuracy (blank) | Notes
Second sheet lists only the questions the platform never answered.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(__file__).with_name("flow_test_results.json")
OUT = Path(__file__).with_name("agent_Flow_Test.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="4F2D7F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")      # no platform response
YELLOW_FILL = PatternFill("solid", fgColor="FFE699")   # inconsistent / partial
GREEN_FILL = PatternFill("solid", fgColor="C6E0B4")    # consistent + data
WRAP = Alignment(wrap_text=True, vertical="top")


def _platform_responded(statuses):
    data = statuses.count("data")
    if data == len(statuses):
        return "Yes"
    if data == 0:
        return "No"
    return "Partial"


def _consistent(runs):
    statuses = [r["class"] for r in runs]
    # status agreement
    if len(set(statuses)) > 1:
        return "No (status varies)"
    # all same status; for data runs, flag big length variation as a soft signal
    if statuses[0] == "data":
        lens = [len(r["response"]) for r in runs]
        if max(lens) and (max(lens) - min(lens)) / max(lens) > 0.5:
            return "Check (length varies)"
        return "Yes"
    if statuses[0] == "empty":
        return "Yes (all empty)"
    return "Yes"


def main():
    data = json.loads(SRC.read_text())

    wb = Workbook()
    ws = wb.active
    ws.title = "Flow Test"

    headers = [
        "SrNo", "Question",
        "Run 1", "Run 2", "Run 3",
        "Platform Responded?", "Consistent?", "MCP Hit", "Avg Latency (s)",
        "Run 1 Response", "Run 2 Response", "Run 3 Response",
        "Raw Data (fill in)", "Accuracy (fill in)", "Notes",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP

    flagged = []
    for entry in sorted(data, key=lambda x: x["srno"]):
        runs = entry["runs"]
        statuses = [r["class"] for r in runs]
        responded = _platform_responded(statuses)
        consistent = _consistent(runs)
        mcp_hits = sum(1 for r in runs if r.get("mcp_hit"))
        lats = [r.get("total_ms") or r.get("wall_ms") or 0 for r in runs]
        avg_lat = round(sum(lats) / len(lats) / 1000, 1) if lats else 0
        resp = [r["response"] for r in runs] + ["", "", ""]

        row = [
            entry["srno"], entry["question"],
            statuses[0] if len(statuses) > 0 else "",
            statuses[1] if len(statuses) > 1 else "",
            statuses[2] if len(statuses) > 2 else "",
            responded, consistent, f"{mcp_hits}/{len(runs)}", avg_lat,
            resp[0], resp[1], resp[2],
            "", "", "",
        ]
        ws.append(row)
        r_idx = ws.max_row
        for c in range(1, len(headers) + 1):
            ws.cell(row=r_idx, column=c).alignment = WRAP
        # row coloring
        fill = None
        if responded == "No":
            fill = RED_FILL
            flagged.append((entry["srno"], entry["question"]))
        elif responded == "Partial" or consistent.startswith(("No", "Check")):
            fill = YELLOW_FILL
        else:
            fill = GREEN_FILL
        for c in (6, 7):  # color the two summary cells
            ws.cell(row=r_idx, column=c).fill = fill

    widths = [6, 46, 9, 9, 9, 16, 18, 9, 13, 48, 48, 48, 30, 14, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # ---- Flagged sheet ----
    ws2 = wb.create_sheet("Flagged - No Platform Data")
    ws2.append(["SrNo", "Question", "Note"])
    for c in range(1, 4):
        ws2.cell(row=1, column=c).fill = HEADER_FILL
        ws2.cell(row=1, column=c).font = HEADER_FONT
    if flagged:
        for sr, q in flagged:
            ws2.append([sr, q, "Platform returned no usable data on all 3 runs"])
    else:
        ws2.append(["—", "None — every question got data on at least one run", ""])
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 60
    ws2.column_dimensions["C"].width = 50
    for row in ws2.iter_rows():
        for cell in row:
            cell.alignment = WRAP

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Flagged (no platform data on any run): {len(flagged)}")
    for sr, q in flagged:
        print(f"  Q{sr}: {q[:70]}")


if __name__ == "__main__":
    main()
