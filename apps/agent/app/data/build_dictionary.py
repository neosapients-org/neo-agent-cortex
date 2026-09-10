"""Compile agent_metrics.xlsx into data_dictionary.json.

The xlsx is the platform's semantic layer (the "data dictionary"). The runtime agent
must NOT depend on openpyxl, so this dev-time script flattens the sheet into a plain
JSON artefact that app/data_dictionary.py loads with the stdlib only.

Run from apps/agent:  python -m app.data.build_dictionary
"""

import json
from pathlib import Path

import openpyxl

_HERE = Path(__file__).parent
_XLSX = _HERE / "agent_metrics.xlsx"
_OUT = _HERE / "data_dictionary.json"

# Header columns in the sheet, in order.
_COLS = ["atom", "name", "display_name", "type", "subtype", "formula", "source_fields", "description"]


def _clean(v) -> str:
    if v is None:
        return ""
    # The sheet contains mojibake em-dashes (â€”) from a bad export; normalise to a dash.
    return str(v).replace("â€”", "—").replace("â€”", "—").strip()


def build() -> list[dict]:
    wb = openpyxl.load_workbook(_XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))[1:]  # skip header
    entries = []
    for r in rows:
        if r[0] is None and r[1] is None:
            continue
        entry = {c: _clean(r[i]) for i, c in enumerate(_COLS)}
        # Drop the heavy formula/source columns from the runtime artefact — validation
        # only needs the vocabulary + description; keep description for explanations.
        entries.append({
            "atom": entry["atom"],
            "name": entry["name"],
            "display_name": entry["display_name"],
            "type": entry["type"],
            "subtype": entry["subtype"],
            "description": entry["description"],
        })
    return entries


def main() -> None:
    entries = build()
    _OUT.write_text(json.dumps(entries, ensure_ascii=False, indent=2))
    metrics = sum(1 for e in entries if e["type"] in ("Metric", "Aggregation"))
    atoms = len({e["atom"] for e in entries})
    print(f"Wrote {len(entries)} entries ({metrics} metric/aggregation, {atoms} domains) -> {_OUT}")


if __name__ == "__main__":
    main()
